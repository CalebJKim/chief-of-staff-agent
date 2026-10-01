"""Build the three editable DOCX templates for the independent task scenarios.

Run with the document-authoring runtime's Python (python-docx required).
"""
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).with_name("templates")


def document(project, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.85)
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string("111827")
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.12
    doc.styles["Title"].font.size = Pt(26)
    doc.styles["Title"].font.bold = True
    doc.styles["Subtitle"].font.size = Pt(13)
    doc.styles["Subtitle"].font.color.rgb = RGBColor.from_string("475569")
    doc.styles["Heading 1"].font.size = Pt(12)
    doc.styles["Heading 1"].font.bold = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    for name in ("Title", "Subtitle", "Heading 1"):
        doc.styles[name].paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = "PRODUCT DEVELOPMENT  /  INTERNAL"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    doc.add_paragraph(project, "Title")
    doc.add_paragraph(subtitle, "Subtitle")
    footer = section.footer.paragraphs[0]
    footer.text = "Internal working document"
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string("64748B")
    doc.core_properties.author = "Product Development"
    doc.core_properties.title = project + " — " + subtitle
    return doc


def section(doc, title, text=None, bullets=()):
    doc.add_paragraph(title, "Heading 1")
    if text:
        doc.add_paragraph(text)
    for value in bullets:
        doc.add_paragraph(value, "List Bullet")


def build():
    OUT.mkdir(exist_ok=True)
    doc = document("AI for Financial Analysis Assistant", "Project Overview")
    doc.add_paragraph("Product lead: Leah Moreno  |  Stage: internal prototype")
    section(doc, "Why we’re building it", "Finance teams spend too much time finding figures across quarterly reports before they can explain what changed. The assistant brings the relevant figures and source passages together, giving analysts a first draft they can check and refine.")
    section(doc, "What the assistant does", bullets=(
        "Summarizes uploaded company financial reports in plain language.",
        "Compares quarters and highlights changes in revenue, operating costs and margins.",
        "Links every numerical claim to the report, page and source figure so analysts can verify it.",
        "Flags missing or inconsistent data instead of filling gaps with estimates.",
    ))
    section(doc, "Who it’s for and what stays out of scope", "The first users are internal finance analysts preparing quarterly reviews. Analysts approve every summary before sharing it. The prototype does not recommend investments, forecast returns or publish analysis automatically.")
    section(doc, "Where the project stands", "The prototype can read a report, answer questions and draft a comparison. Evaluation uses synthetic company reports. Simple layouts work well, but multi-column tables sometimes mix periods or lose units. Citations also need checking after extraction. These issues must be resolved before the finance-team review.")
    section(doc, "Leah’s role", "Leah is taking over product coordination: keep the scope focused, agree on acceptance criteria with Engineering and organize feedback from the finance team. Start with Progress and Findings for the current evidence, then Next Steps for the work to coordinate.")
    doc.save(OUT / "financial-analysis-overview.docx")

    doc = document("AI for Financial Analysis Assistant", "Next Steps")
    doc.add_paragraph("Product lead: Leah Moreno  |  Goal: a reliable internal review")
    section(doc, "Start here", "Read the Project Overview, then the four-slide Progress and Findings deck. Meet with Engineering to reproduce the extraction and citation problems before agreeing on the next prototype build.")
    section(doc, "1. Fix financial table extraction", "Engineering should preserve period headers, row labels, units and footnotes when reading multi-column tables. Keep the failing examples as regression cases. A comparison is ready for review only when both periods map to the correct source cells.")
    section(doc, "2. Verify numerical claims and citations", "Check each generated figure against its cited page and table cell. Test missing pages, contradictory figures and ambiguous units. If a figure cannot be verified, the assistant should flag it and leave the comparison incomplete.")
    section(doc, "3. Arrange the finance-team review", "Leah should identify reviewers and agree on a small set of realistic quarterly-review questions. Use synthetic reports for the first session. Ask reviewers whether the draft saves time, whether the citations are easy to follow and where the wording could mislead.")
    section(doc, "Decisions to close before scheduling", bullets=(
        "Which report layouts and financial metrics are in the first test set?",
        "What extraction and citation checks must pass before reviewers use the build?",
        "Who will review the results, and when can Engineering provide the corrected prototype?",
    ))
    section(doc, "Expected handoff from Leah", "A short review plan with agreed checks, named owners and a proposed review date. No review date or release commitment has been agreed yet.")
    doc.save(OUT / "financial-analysis-next-steps.docx")

    build_notes()


def build_notes():
    """Build the open design outline without regenerating unrelated templates."""
    OUT.mkdir(exist_ok=True)
    doc = document("Local AI Meeting Notes Assistant", "Design Outline")
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = doc.styles["Title"].element.find(qn("w:pPr"))
    if title_properties is not None:
        for border in list(title_properties.findall(qn("w:pBdr"))):
            title_properties.remove(border)
    for run in doc.sections[0].header.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph("Engineering lead: Evan Mercer  |  Stage: design definition")
    section(doc, "Purpose", "Help people capture useful meeting notes without sending meeting recordings or transcripts to a cloud model.")
    section(doc, "Proposed design")
    for text in (
        "Diagram of the user experience before, during, and after a meeting. Decide whether notes appear live or after the meeting.",
        "Outline how the assistant turns meeting audio into useful notes while keeping processing on the device.",
        "Decide how users review and correct summaries, action items, and deadlines before sharing.",
        "Choose what the first prototype should include, what can wait, and explain the main tradeoffs.",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(14)
        p.add_run("TODO: ").bold = True
        p.add_run(text)
    doc.save(OUT / "local-meeting-notes-overview.docx")


if __name__ == "__main__":
    build()
